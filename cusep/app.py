import csv
import io
import json
import os
import secrets
import sqlite3
import tempfile
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file, session
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .crawl import Jobs, crawl_company, import_catalogue, import_directory
from .db import Database, now
from .extract import SECTORS, canonical, classify, document_candidates, ocr_image, spreadsheet_rows

DEFAULT_SOURCE = 'https://www.baskentosb.org/tr/firma-listesi/'


def create_app(data_dir=None):
    app = Flask(__name__)
    directory = Path(data_dir or os.environ.get('CUSEP_DATA_DIR', Path(__file__).resolve().parent.parent / 'data'))
    directory.mkdir(parents=True, exist_ok=True)
    secret_path = directory / '.session-key'
    if not secret_path.exists():
        with secret_path.open('x') as f: f.write(secrets.token_hex(32))
        secret_path.chmod(0o600)
    app.secret_key = secret_path.read_text().strip()
    app.config.update(MAX_CONTENT_LENGTH=20*1024*1024, SESSION_COOKIE_SAMESITE='Strict', SESSION_COOKIE_HTTPONLY=True)
    db = Database(directory / 'cusep.sqlite3')
    jobs = Jobs(db)
    app.extensions['db'] = db
    app.extensions['jobs'] = jobs

    @app.before_request
    def csrf():
        if request.method in ('POST', 'PATCH', 'DELETE'):
            token = session.get('csrf')
            if not token or not secrets.compare_digest(request.headers.get('X-CSRF-Token', ''), token):
                return jsonify(error='Session expired. Reload the page and try again.'), 403

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['X-Frame-Options'] = 'DENY'
        response.headers['Referrer-Policy'] = 'no-referrer'
        response.headers['Content-Security-Policy'] = "default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        return response

    @app.errorhandler(ValueError)
    def bad_input(error):
        return jsonify(error=str(error)), 400

    @app.errorhandler(413)
    def too_large(error):
        return jsonify(error='File exceeds the 20 MB upload limit.'), 413

    @app.get('/')
    def home():
        session.setdefault('csrf', secrets.token_hex(32))
        return render_template('index.html', csrf=session['csrf'], source=DEFAULT_SOURCE)

    @app.get('/api/stats')
    def stats():
        return jsonify(
            companies=db.query('SELECT COUNT(*) n FROM companies')[0]['n'],
            listings=db.query('SELECT COUNT(*) n FROM listings')[0]['n'],
            source_breakdown=db.query("SELECT source_sector sector,COUNT(*) n FROM companies WHERE source_sector<>'' GROUP BY source_sector ORDER BY n DESC"),
            products=db.query('SELECT COUNT(*) n FROM products')[0]['n'],
            documents=db.query('SELECT COUNT(*) n FROM documents')[0]['n'],
            needs_review=db.query("SELECT (SELECT COUNT(*) FROM companies WHERE status='needs_review') + (SELECT COUNT(*) FROM products WHERE status='needs_review') n")[0]['n'],
            sectors=list(SECTORS),
            breakdown=db.query("SELECT j.value sector,COUNT(*) n FROM companies c,json_each(c.sectors) j GROUP BY j.value ORDER BY n DESC"),
            recent_jobs=db.query('SELECT * FROM jobs ORDER BY id DESC LIMIT 5'))

    def filters(product=False):
        q = request.args.get('q', '').strip()
        sector = request.args.get('sector', '')
        status = request.args.get('status', '')
        clauses, args = [], []
        if q:
            fields = "c.name || ' ' || p.name || ' ' || p.specs" if product else "c.name || ' ' || c.description || ' ' || c.website"
            clauses.append(f'fold({fields}) LIKE fold(?)')
            args.append('%' + q + '%')
        if sector:
            if sector.startswith('source:'):
                clauses.append('c.source_sector=?'); args.append(sector[7:])
            else:
                clauses.append('EXISTS (SELECT 1 FROM json_each(c.sectors) WHERE value=?)'); args.append(sector)
        if status:
            clauses.append(('p' if product else 'c') + '.status=?')
            args.append(status)
        if product and request.args.get('company_id'):
            clauses.append('p.company_id=?'); args.append(int(request.args['company_id']))
        return (' WHERE ' + ' AND '.join(clauses)) if clauses else '', args

    @app.get('/api/companies')
    def companies():
        where, args = filters()
        page = max(1, int(request.args.get('page', 1)))
        rows = db.query('SELECT c.*, (SELECT COUNT(*) FROM products p WHERE p.company_id=c.id) product_count FROM companies c' + where + ' ORDER BY c.name COLLATE NOCASE LIMIT 50 OFFSET ?', args + [(page-1)*50])
        total = db.query('SELECT COUNT(*) n FROM companies c' + where, args)[0]['n']
        return jsonify(items=rows, total=total, page=page)

    @app.get('/api/company-options')
    def company_options():
        return jsonify(items=db.query('SELECT id,name,website FROM companies ORDER BY name LIMIT 10000'))

    @app.get('/api/companies/<int:cid>')
    def company(cid):
        rows = db.query('SELECT * FROM companies WHERE id=?', (cid,))
        if not rows: return jsonify(error='Company not found'), 404
        return jsonify(company=rows[0], listings=db.query('SELECT * FROM listings WHERE company_id=? ORDER BY address', (cid,)), products=db.query('SELECT * FROM products WHERE company_id=? ORDER BY name', (cid,)), documents=db.query('SELECT id,title,source_url,kind,fetched_at,sha256 FROM documents WHERE company_id=? ORDER BY fetched_at DESC,id DESC', (cid,)))

    @app.post('/api/companies')
    def add_company():
        data = request.get_json()
        website = safe_website(data.get('website', ''))
        cid = db.company(data.get('name', ''), website=website, description=data.get('description', ''), sectors=valid_sectors(data.get('sectors', [])), evidence='Manually entered; verify against the company source.')
        return jsonify(id=cid), 201

    @app.patch('/api/companies/<int:cid>')
    def update_company(cid):
        if not db.query('SELECT id FROM companies WHERE id=?', (cid,)): return jsonify(error='Company not found'), 404
        data = request.get_json()
        sectors = valid_sectors(data.get('sectors', []))
        status = valid_status(data.get('status', 'needs_review'))
        # Identity is deliberately retained so a directory refresh still matches the imported company.
        name = str(data.get('name', '')).strip()
        if not name: raise ValueError('Company name is required.')
        db.execute('UPDATE companies SET name=?,website=?,sectors=?,description=?,status=?,locked=?,updated_at=? WHERE id=?',
                   (name[:500], safe_website(data.get('website', '')), json.dumps(sectors, ensure_ascii=False), str(data.get('description', ''))[:12000], status, int(bool(data.get('locked', True))), now(), cid))
        return jsonify(ok=True)

    @app.get('/api/products')
    def products():
        where, args = filters(True)
        page = max(1, int(request.args.get('page', 1)))
        sql = ' FROM products p JOIN companies c ON c.id=p.company_id'
        return jsonify(items=db.query('SELECT p.*,c.name company_name,c.sectors' + sql + where + ' ORDER BY p.updated_at DESC LIMIT 50 OFFSET ?', args+[(page-1)*50]),
                       total=db.query('SELECT COUNT(*) n' + sql + where, args)[0]['n'], page=page)

    @app.post('/api/products')
    def add_product():
        data = request.get_json()
        cid = int(data.get('company_id', 0))
        if not db.query('SELECT id FROM companies WHERE id=?', (cid,)): raise ValueError('Select a company.')
        name = str(data.get('name', '')).strip()
        if not name: raise ValueError('Product name is required.')
        specs = valid_specs(data.get('specs', {}))
        reference = str(data.get('source_url', '')).strip()
        document_id = None
        if reference.startswith('upload:'):
            documents = db.query('SELECT id FROM documents WHERE source_url=? AND company_id=?', (reference.split('#')[0],cid))
            if not documents: raise ValueError('Uploaded source identifier does not belong to this company.')
            source = reference
            document_id = documents[0]['id']
        else:
            source = safe_website(reference)
        if not source: raise ValueError('A product source URL is required. For uploaded catalogues, select a document and use its source identifier.')
        pid = db.product(cid, name, specs, source, str(data.get('evidence', 'Manual entry; verify with the manufacturer.')), document_id)
        return jsonify(id=pid), 201

    @app.patch('/api/products/<int:pid>')
    def update_product(pid):
        if not db.query('SELECT id FROM products WHERE id=?', (pid,)): return jsonify(error='Product not found'), 404
        data = request.get_json()
        name = str(data.get('name', '')).strip()
        if not name: raise ValueError('Product name is required.')
        db.execute('UPDATE products SET name=?,specs=?,status=?,locked=?,updated_at=? WHERE id=?',
                   (name[:500], json.dumps(valid_specs(data.get('specs', {})), ensure_ascii=False), valid_status(data.get('status', 'needs_review')), int(bool(data.get('locked', True))), now(), pid))
        return jsonify(ok=True)

    @app.get('/api/documents')
    def documents():
        q = request.args.get('q', '')
        rows = db.query('''SELECT d.id,d.company_id,d.source_url,d.kind,d.title,d.fetched_at,d.sha256,c.name company_name
         FROM documents d LEFT JOIN companies c ON c.id=d.company_id
         WHERE fold(d.title || ' ' || d.content) LIKE fold(?) ORDER BY d.fetched_at DESC,d.id DESC LIMIT 100''', ('%'+q+'%',))
        return jsonify(items=rows)

    @app.get('/api/documents/<int:did>')
    def document(did):
        rows = db.query('SELECT id,company_id,source_url,kind,title,sha256,content,fetched_at FROM documents WHERE id=?', (did,))
        if not rows: return jsonify(error='Source not found'), 404
        return jsonify(rows[0])

    @app.get('/api/documents/<int:did>/file')
    def original_file(did):
        rows = db.query('SELECT raw,title FROM documents WHERE id=?', (did,))
        if not rows or rows[0]['raw'] is None: return jsonify(error='Original bytes unavailable'), 404
        name = Path(rows[0]['title'].replace('\\', '/')).name or 'original-document'
        return send_file(io.BytesIO(rows[0]['raw']), mimetype='application/octet-stream', as_attachment=True, download_name=name)

    @app.post('/api/import/directory')
    def directory_import():
        data = request.get_json()
        url = canonical(data.get('url', DEFAULT_SOURCE))
        if not url: raise ValueError('Enter a valid directory URL.')
        return jsonify(job_id=jobs.start('directory', lambda log: import_directory(db, url, log)))

    @app.post('/api/import/company')
    def company_import():
        data = request.get_json()
        ids = data.get('company_ids')
        if not ids or not isinstance(ids, list) or len(ids) > 1000: raise ValueError('Select between 1 and 1,000 companies.')
        ids = list(dict.fromkeys(int(x) for x in ids))
        limit = int(data.get('max_pages', 30))
        if not 1 <= limit <= 200: raise ValueError('Page limit must be between 1 and 200 per company.')
        def run(log):
            failed = 0
            for cid in ids:
                try: log(crawl_company(db, cid, log, max_pages=limit))
                except ValueError as exc: failed += 1; log(str(exc))
            if failed: raise ValueError(f'{failed}/{len(ids)} company crawls had issues. Successfully extracted records were retained; see activity details.')
            return f'Finished {len(ids)} company website crawls. Review extracted candidates.'
        return jsonify(job_id=jobs.start('company websites', run))

    @app.post('/api/import/file')
    def file_import():
        upload = request.files.get('file')
        if not upload or not upload.filename: raise ValueError('Select a file.')
        filename = Path(upload.filename.replace('\\', '/')).name
        ext = Path(filename).suffix.lower()
        if ext not in ('.csv', '.xlsx', '.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp'):
            raise ValueError('Supported formats: CSV, XLSX, PDF, PNG, JPEG, TIFF, WebP.')
        raw = upload.read()
        if not raw: raise ValueError('File is empty.')
        cid = request.form.get('company_id', '').strip()
        source = canonical(request.form.get('source_url', ''))
        import hashlib
        identity = 'upload:' + hashlib.sha256(f'{cid or "spreadsheet"}:{filename.casefold()}'.encode()).hexdigest()
        if ext in ('.csv', '.xlsx'):
            rows = list(spreadsheet_rows(raw, filename))  # validate before starting; never execute formulas
            def run(log):
                count = 0
                doc = db.document(identity, json.dumps(rows,ensure_ascii=False,indent=2), 'spreadsheet', filename, raw=raw)
                for row in rows:
                    sectors = [x.strip() for x in row['sector'].split(';') if x.strip()]
                    c = db.company(row['company'], safe_website(row['website']), sectors=sectors,
                                   directory_url=canonical(row.get('directory_url', '')), source_sector=row.get('source_sector',''), description=row['description'], evidence=f'Spreadsheet {filename}, row {row["row"]}')
                    if row['product']:
                        db.product(c, row['product'], row['specs'], canonical(row['source_url']) or source or identity,
                                   f'{filename} row {row["row"]}', doc)
                    count += 1
                    if count % 50 == 0: log(f'Imported {count}/{len(rows)} spreadsheet rows')
                return f'Imported/refreshed {count} spreadsheet rows. Review source attribution and sector labels.'
        else:
            if not cid or not db.query('SELECT id FROM companies WHERE id=?', (cid,)):
                raise ValueError('Select the company that owns this catalogue/image before uploading.')
            cid = int(cid)
            def run(log):
                log('Extracting ' + filename)
                if ext == '.pdf':
                    count, warnings = import_catalogue(db, raw, source or identity, filename, cid, log)
                    return f'Stored catalogue; {count} product candidates extracted. {len(warnings)} OCR warnings; inspect source pages.'
                text = ocr_image(raw)
                if not text.strip(): raise ValueError('OCR found no readable text. Try a clearer image.')
                doc = db.document(source or identity, text, 'image OCR', filename, cid, raw=raw)
                products = document_candidates(text)
                for p in products: db.product(cid, p['name'], p['specs'], source or identity, p['evidence'], doc)
                return f'Stored OCR text; {len(products)} product candidates. Verify names, numbers and units against the image.'
        return jsonify(job_id=jobs.start('file: ' + filename, run))

    @app.get('/api/jobs')
    def activity():
        return jsonify(items=db.query('SELECT * FROM jobs ORDER BY id DESC LIMIT 100'))

    @app.post('/api/jobs/<int:jid>/cancel')
    def cancel(jid):
        db.execute("UPDATE jobs SET cancel_requested=1 WHERE id=? AND state IN ('queued','running')", (jid,))
        return jsonify(ok=True)

    @app.get('/api/export/<kind>')
    def export(kind):
        companies = db.query('SELECT name,website,source_sector,sectors,description,directory_url,status,updated_at,last_seen FROM companies ORDER BY source_sector,name')
        listings = db.query('SELECT name,address,phone,email,source_sector,operations,website,source_url,last_seen FROM listings ORDER BY source_sector,name,address')
        products = db.query('''SELECT c.name company,p.name product,p.specs,p.source_url,p.evidence,p.status,p.updated_at,p.last_seen
                              FROM products p JOIN companies c ON c.id=p.company_id ORDER BY c.name,p.name''')
        for c in companies: c['sectors'] = '; '.join(json.loads(c['sectors']))
        if kind == 'backup':
            # SQLite backup API includes committed WAL data consistently.
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / 'cusep.sqlite3'
                with db.connect() as src, sqlite3.connect(path) as dest:
                    src.backup(dest)
                    dest.execute('PRAGMA journal_mode=DELETE')
                raw = path.read_bytes()
            return send_file(io.BytesIO(raw), mimetype='application/vnd.sqlite3', as_attachment=True, download_name='cusep-backup.sqlite3')
        if kind == 'xlsx':
            book = Workbook(); book.remove(book.active)
            for title, records, columns in [('Companies', companies, ['name','website','source_sector','sectors','description','directory_url','status','updated_at','last_seen']),
                                             ('Locations', listings, ['name','address','phone','email','source_sector','operations','website','source_url','last_seen']),
                                             ('Products', products, ['company','product','specs','source_url','evidence','status','updated_at','last_seen'])]:
                sheet = book.create_sheet(title)
                sheet.append(columns)
                for row in records: sheet.append([excel_safe(row.get(k, '')) for k in columns])
                sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
                for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width = 28
            spec_sheet = book.create_sheet('Specifications'); spec_sheet.append(['company','product','specification','value','source_url'])
            for p in products:
                for k, v in json.loads(p['specs']).items(): spec_sheet.append([excel_safe(x) for x in (p['company'],p['product'],k,v,p['source_url'])])
            output = io.BytesIO(); book.save(output); output.seek(0)
            return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='cusep-database.xlsx')
        if kind not in ('companies.csv', 'products.csv', 'listings.csv'): raise ValueError('Choose xlsx, companies.csv, products.csv, or backup.')
        records = companies if kind.startswith('companies') else listings if kind.startswith('listings') else products
        columns = ['name','website','source_sector','sectors','description','directory_url','status','updated_at','last_seen'] if kind.startswith('companies') else ['company','product','specs','source_url','evidence','status','updated_at','last_seen']
        if kind == 'listings.csv': columns = ['name','address','phone','email','source_sector','operations','website','source_url','last_seen']
        output = io.StringIO(); writer = csv.writer(output); writer.writerow(columns)
        for row in records: writer.writerow([excel_safe(row.get(k,'')) for k in columns])
        return send_file(io.BytesIO(output.getvalue().encode('utf-8-sig')), mimetype='text/csv', as_attachment=True, download_name=kind)

    @app.get('/health')
    def health():
        db.query('SELECT COUNT(*) FROM companies')
        return jsonify(status='ok')

    return app


def safe_website(value):
    if not str(value).strip(): return ''
    result = canonical(str(value))
    if not result: raise ValueError('Website/source must be an absolute HTTP(S) URL.')
    return result


def valid_status(value):
    if value not in ('needs_review', 'reviewed', 'rejected'): raise ValueError('Invalid review status.')
    return value


def valid_sectors(value):
    if not isinstance(value, list) or len(value) > 30 or any(not isinstance(x, str) or len(x) > 150 for x in value):
        raise ValueError('Sectors must be a list of labels (up to 30).')
    return [s.strip() for s in value if s.strip()]


def valid_specs(value):
    if not isinstance(value, dict) or len(value) > 100 or any(not isinstance(k,str) or not isinstance(v,(str,int,float)) for k,v in value.items()):
        raise ValueError('Specifications must be a JSON object of up to 100 names and text/numeric values.')
    if len(json.dumps(value)) > 30000: raise ValueError('Specifications are too large (30,000 characters maximum).')
    return value


def excel_safe(value):
    value = str(value)
    if value.lstrip().startswith(('=', '+', '-', '@')): value = "'" + value
    return ILLEGAL_CHARACTERS_RE.sub('',value)[:32000]
